import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tablica import Tablica


def vybrat_nomer(kolichestvo, tekst, mozhno_nazad=True):
    vybor = input(tekst)
    while not (
        mozhno_nazad and vybor == "-1"
    ) and (not vybor.isdigit() or not 1 <= int(vybor) <= kolichestvo):
        vybor = input("Нет такого номера. Введите еще раз: ")
    if mozhno_nazad and vybor == "-1":
        return -1
    return int(vybor)


def vybrat_neskolko_tablic(kolichestvo):
    vybor = input("Введите номера таблиц через запятую: ").strip()
    while True:
        if vybor == "-1":
            return None
        try:
            nomera = [int(nomer.strip()) for nomer in vybor.split(",")]
            nomera = list(dict.fromkeys(nomera))
            if nomera and all(1 <= nomer <= kolichestvo for nomer in nomera):
                return nomera
        except ValueError:
            pass
        vybor = input("Нет таких номеров. Введите еще раз: ").strip()


def sohranit_excel(tablica):
    imya_bez_rasshireniya = os.path.splitext(tablica.imya)[0]
    tablica.imya = imya_bez_rasshireniya + ".xlsx"
    tablica.df.to_excel(tablica.imya, index=False)

    # Для ручной проверки даём простой выбор "нет/да". Значение "да"
    # подсвечивается зелёным и хорошо видно даже в старых версиях Excel.
    if "проверено" not in tablica.df.columns:
        return

    kniga = load_workbook(tablica.imya)
    list_excel = kniga.active
    nomer_kolonki = list(tablica.df.columns).index("проверено") + 1
    bukva_kolonki = get_column_letter(nomer_kolonki)

    if list_excel.max_row >= 2:
        diapazon = f"{bukva_kolonki}2:{bukva_kolonki}{list_excel.max_row}"
        vybor_da_net = DataValidation(
            type="list",
            formula1='"нет,да"',
            allow_blank=False,
        )
        vybor_da_net.error = "Выберите 'да' или 'нет'"
        vybor_da_net.errorTitle = "Неверное значение"
        vybor_da_net.prompt = "Выберите 'да', если операция проверена"
        vybor_da_net.promptTitle = "Проверка операции"
        vybor_da_net.showErrorMessage = True
        vybor_da_net.showInputMessage = True
        list_excel.add_data_validation(vybor_da_net)
        vybor_da_net.add(diapazon)

        zelenaya_zalivka = PatternFill(
            fill_type="solid", fgColor="C6EFCE"
        )
        list_excel.conditional_formatting.add(
            diapazon,
            FormulaRule(
                formula=[f'{bukva_kolonki}2="да"'],
                fill=zelenaya_zalivka,
            ),
        )

    kniga.save(tablica.imya)


def sobrat_operacii_ispolnitelya(tablicy, papka_goda=".", god=""):
    istochniki = [
        tablica for tablica in tablicy if "Исполнитель" in tablica.df.columns
    ]

    if not istochniki:
        print("В текущей сессии нет таблиц с колонкой 'Исполнитель'")
        return None, False

    print("\nВыберите таблицы с операциями:")
    for nomer, tablica in enumerate(istochniki, start=1):
        print(nomer, "-", tablica.imya)
    print("-1 - вернуться в главное меню")

    nomera_istochnikov = vybrat_neskolko_tablic(len(istochniki))
    if nomera_istochnikov is None:
        return None, False
    vybrannye_istochniki = [
        istochniki[nomer - 1] for nomer in nomera_istochnikov
    ]

    lyudi = {
        1: "Алексей",
        2: "Владимир",
        3: "Дмитрий",
        4: "Андрей",
    }

    print("\nИсполнители:")
    for nomer, chelovek in lyudi.items():
        print(nomer, "-", chelovek)
    print("-1 - вернуться в главное меню")

    nomer_cheloveka = vybrat_nomer(4, "Введите номер исполнителя: ")
    if nomer_cheloveka == -1:
        return None, False
    chelovek = lyudi[nomer_cheloveka]

    operacii_iz_tablic = []
    for istochnik in vybrannye_istochniki:
        operacii_iz_tablic.append(
            istochnik.df[
                istochnik.df["Исполнитель"]
                .astype(str)
                .str.strip()
                .str.lower()
                == chelovek.lower()
            ].copy()
        )
    operacii = pd.concat(operacii_iz_tablic, ignore_index=True)

    imya_rezultata = os.path.join(papka_goda, f"{chelovek}_{god}.xlsx")

    if os.path.exists(imya_rezultata):
        print("\nНайден файл:", imya_rezultata)
        print("1 - дописать новые операции в конец")
        print("2 - ничего не делать")
        print("-1 - вернуться в главное меню")
        deystvie = vybrat_nomer(2, "Введите номер пункта: ")

        if deystvie in [-1, 2]:
            print("Сохранение отменено")
            return None, False

        try:
            staraya_tablica = pd.read_excel(imya_rezultata)
        except Exception as oshibka:
            print("Ошибка чтения файла:", oshibka)
            return None, False
    else:
        print("\nФайл", imya_rezultata, "пока не существует")
        print("1 - создать новую таблицу")
        print("2 - считать старую таблицу и дополнить ее")
        print("-1 - вернуться в главное меню")
        deystvie = vybrat_nomer(2, "Введите номер пункта: ")

        if deystvie == -1:
            return None, False

        if deystvie == 1:
            staraya_tablica = pd.DataFrame()
        else:
            imya_starogo = input("Введите название старого файла (-1 — назад): ")
            if imya_starogo.strip() == "-1":
                return None, False
            try:
                staraya_tablica = pd.read_excel(imya_starogo)
            except Exception as oshibka:
                print("Ошибка чтения файла:", oshibka)
                return None, False

    # Все новые операции ещё не проверены. В старом файле сохраняем уже
    # проставленные ответы, а пустые значения считаем непроверенными.
    operacii["проверено"] = "нет"
    if "проверено" not in staraya_tablica.columns:
        staraya_tablica["проверено"] = "нет"
    else:
        provereno = staraya_tablica["проверено"].fillna("").astype(str).str.strip()
        staraya_tablica.loc[provereno == "", "проверено"] = "нет"

    novyi_df = pd.concat(
        [staraya_tablica, operacii], ignore_index=True
    )
    # Колонка проверки всегда должна быть последней.
    ostalnye_kolonki = [
        kolonka for kolonka in novyi_df.columns if kolonka != "проверено"
    ]
    novyi_df = novyi_df[ostalnye_kolonki + ["проверено"]]

    # Если итоговый файл уже был загружен в этой сессии, обновляем его DataFrame.
    rezultat = None
    for tablica in tablicy:
        if os.path.abspath(tablica.imya) == os.path.abspath(imya_rezultata):
            rezultat = tablica
            rezultat.df = novyi_df
            break

    if rezultat is None:
        rezultat = Tablica(imya_rezultata, novyi_df)
        dobavit_v_sessiyu = True
    else:
        dobavit_v_sessiyu = False

    sohranit_excel(rezultat)
    print("Готово. Операции добавлены в файл:", rezultat.imya)
    return rezultat, dobavit_v_sessiyu
