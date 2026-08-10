import os
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tablica import Tablica


KOLONKI = [
    "Дата",
    "Фирма",
    "Контрагент",
    "Пришло/ушло",
    "Тип операции",
    "Наименование",
    "Сумма",
    "Комментарий",
]

KOMPANII = [
    ("Альтэгра", ["альтэгра", "альтегра"]),
    ("АВК", ["авк"]),
    ("Билд", ["билд"]),
    ("Вектор", ["вектор"]),
    ("Макрон", ["макрон"]),
    ("Позитрон", ["позитрон"]),
    ("Сити", ["сити"]),
    ("Кит", ["кит"]),
    ("Энергопоинт", ["энергопоинт", "энерго поинт"]),
    ("Факторион", ["факторион"]),
]


def normalizovat_tekst(znachenie):
    if pd.isna(znachenie):
        return ""
    return " ".join(
        str(znachenie).strip().lower().replace("ё", "е").split()
    )


def nazvanie_firmy(imya_fayla):
    osnova = os.path.splitext(os.path.basename(imya_fayla))[0]
    return osnova.split("_", 1)[0].strip()


def nayti_nashu_kompaniyu(znachenie):
    tekst = normalizovat_tekst(znachenie)
    if not tekst:
        return None
    for nazvanie, varianty in KOMPANII:
        for variant in varianty:
            shablon = rf"(?<![а-яa-z0-9]){re.escape(variant)}(?![а-яa-z0-9])"
            if re.search(shablon, tekst, flags=re.IGNORECASE):
                return nazvanie
    if re.search(r"(^|\W)эп($|\W)", tekst, flags=re.IGNORECASE):
        return "Энергопоинт"
    return None


def denezhnoe_chislo(znachenie):
    if pd.isna(znachenie) or str(znachenie).strip() == "":
        return Decimal("0.00")
    try:
        tekst = str(znachenie).replace(" ", "").replace("\u00a0", "")
        tekst = tekst.replace(",", ".")
        return Decimal(tekst).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
    except (InvalidOperation, ValueError, TypeError):
        return None


def preobrazovat_datu(znachenie):
    if pd.isna(znachenie) or str(znachenie).strip() == "":
        return pd.NaT
    data = pd.to_datetime(znachenie, dayfirst=True, errors="coerce")
    if pd.isna(data):
        return pd.NaT
    return pd.Timestamp(data).normalize()


def oformit_tablicu_vseh_operaciy(imya_fayla):
    kniga = load_workbook(imya_fayla)
    list_excel = kniga.active
    list_excel.freeze_panes = "A2"
    list_excel.auto_filter.ref = list_excel.dimensions

    zalivka = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for yacheyka in list_excel[1]:
        yacheyka.font = Font(bold=True)
        yacheyka.fill = zalivka

    shiriny = [14, 20, 30, 15, 20, 32, 18, 55]
    for nomer, shirina in enumerate(shiriny, start=1):
        list_excel.column_dimensions[get_column_letter(nomer)].width = shirina
    for yacheyka in list_excel["A"][1:]:
        yacheyka.number_format = "dd.mm.yyyy"
    for yacheyka in list_excel["G"][1:]:
        yacheyka.number_format = "#,##0.00"
    kniga.save(imya_fayla)


def vybrat_istochniki(tablicy):
    obyazatelnye = {
        "Дата операции",
        "Корреспондент",
        "Оборот Дт",
        "Оборот Кт",
        "Тип операции",
    }
    istochniki = [
        tablica for tablica in tablicy
        if obyazatelnye.issubset(tablica.df.columns)
    ]
    if not istochniki:
        print(
            "В текущей сессии нет подходящих таблиц. "
            "Сначала добавьте типы операций."
        )
        return None

    print("\nВыберите таблицы компаний:")
    for nomer, tablica in enumerate(istochniki, start=1):
        print(nomer, "-", tablica.imya)
    print("-1 - вернуться назад")

    vybor = input("Введите номера таблиц через запятую: ").strip()
    while True:
        if vybor == "-1":
            return None
        try:
            nomera = [int(nomer.strip()) for nomer in vybor.split(",")]
            nomera = list(dict.fromkeys(nomera))
            if nomera and all(1 <= nomer <= len(istochniki) for nomer in nomera):
                vybrannye = [istochniki[nomer - 1] for nomer in nomera]
                break
        except ValueError:
            pass
        vybor = input("Нет таких номеров. Введите еще раз: ").strip()

    identifikatory = []
    for tablica in vybrannye:
        imya = nazvanie_firmy(tablica.imya)
        identifikatory.append(
            normalizovat_tekst(nayti_nashu_kompaniyu(imya) or imya)
        )
    if len(identifikatory) != len(set(identifikatory)):
        print(
            "Ошибка: выбрано несколько таблиц одной фирмы. "
            "Выберите только один готовый файл для каждой фирмы."
        )
        return None
    return vybrannye


def dannye_operacii(stroka):
    naimenovanie = stroka.get("Наименование")
    if pd.isna(naimenovanie):
        naimenovanie = ""
    kommentariy = stroka.get("Назначение платежа")
    if pd.isna(kommentariy):
        kommentariy = ""
    tip = stroka.get("Тип операции")
    if pd.isna(tip):
        tip = ""
    return str(tip), str(naimenovanie), str(kommentariy)


def sozdat_stroku(data, firma, kontragent, napravlenie, operaciya, summa):
    tip, naimenovanie, kommentariy = dannye_operacii(operaciya)
    return {
        "Дата": data,
        "Фирма": firma,
        "Контрагент": kontragent,
        "Пришло/ушло": napravlenie,
        "Тип операции": tip,
        "Наименование": naimenovanie,
        "Сумма": float(summa),
        "Комментарий": kommentariy,
    }


def kluch_daty(data):
    if pd.isna(data):
        return ""
    return data.strftime("%Y-%m-%d")


def sostavit_avtomaticheski(vybrannye):
    obychnye_stroki = []
    vnutrennie_perevody = defaultdict(lambda: defaultdict(list))
    propusheno = 0

    vidimye_imena = {}
    for tablica in vybrannye:
        imya_iz_fayla = nazvanie_firmy(tablica.imya)
        kanonicheskoe = nayti_nashu_kompaniyu(imya_iz_fayla)
        identifikator = normalizovat_tekst(kanonicheskoe or imya_iz_fayla)
        vidimye_imena[identifikator] = kanonicheskoe or imya_iz_fayla

    for tablica in vybrannye:
        imya_iz_fayla = nazvanie_firmy(tablica.imya)
        nasha_firma = nayti_nashu_kompaniyu(imya_iz_fayla) or imya_iz_fayla
        id_firmy = normalizovat_tekst(nasha_firma)

        for _, stroka in tablica.df.iterrows():
            data = preobrazovat_datu(stroka.get("Дата операции"))
            dt = denezhnoe_chislo(stroka.get("Оборот Дт"))
            kt = denezhnoe_chislo(stroka.get("Оборот Кт"))
            if dt is None or kt is None:
                propusheno += 1
                continue
            if dt == 0 and kt == 0:
                continue

            kontragent_tekst = stroka.get("Корреспондент")
            if pd.isna(kontragent_tekst):
                kontragent_tekst = ""
            kontragent_tekst = str(kontragent_tekst)
            nash_kontragent = nayti_nashu_kompaniyu(kontragent_tekst)

            storony = []
            if dt != 0:
                storony.append(("ушло", dt))
            if kt != 0:
                storony.append(("пришло", kt))

            for napravlenie, summa in storony:
                if nash_kontragent is None:
                    obychnye_stroki.append(sozdat_stroku(
                        data,
                        nasha_firma,
                        kontragent_tekst,
                        napravlenie,
                        stroka,
                        summa,
                    ))
                    continue

                id_kontragenta = normalizovat_tekst(nash_kontragent)
                vidimye_imena.setdefault(id_kontragenta, nash_kontragent)
                if napravlenie == "ушло":
                    otpravitel, poluchatel = id_firmy, id_kontragenta
                else:
                    otpravitel, poluchatel = id_kontragenta, id_firmy

                kluch = (
                    kluch_daty(data),
                    otpravitel,
                    poluchatel,
                    summa,
                )
                vnutrennie_perevody[kluch][id_firmy].append({
                    "Дата": data,
                    "Строка": stroka,
                })

    for kluch, po_istochnikam in vnutrennie_perevody.items():
        _, otpravitel, poluchatel, summa = kluch
        kolichestvo = max(len(stroki) for stroki in po_istochnikam.values())
        stroki_otpravitelya = po_istochnikam.get(otpravitel, [])
        stroki_poluchatelya = po_istochnikam.get(poluchatel, [])

        for nomer in range(kolichestvo):
            operaciya_otpravitelya = (
                stroki_otpravitelya[nomer]
                if nomer < len(stroki_otpravitelya)
                else None
            )
            operaciya_poluchatelya = (
                stroki_poluchatelya[nomer]
                if nomer < len(stroki_poluchatelya)
                else None
            )
            obrazec = operaciya_otpravitelya or operaciya_poluchatelya
            data = obrazec["Дата"]

            istochnik_otpravitelya = (
                operaciya_otpravitelya or operaciya_poluchatelya
            )["Строка"]
            istochnik_poluchatelya = (
                operaciya_poluchatelya or operaciya_otpravitelya
            )["Строка"]
            imya_otpravitelya = vidimye_imena.get(otpravitel, otpravitel)
            imya_poluchatelya = vidimye_imena.get(poluchatel, poluchatel)

            obychnye_stroki.append(sozdat_stroku(
                data,
                imya_otpravitelya,
                imya_poluchatelya,
                "ушло",
                istochnik_otpravitelya,
                summa,
            ))
            obychnye_stroki.append(sozdat_stroku(
                data,
                imya_poluchatelya,
                imya_otpravitelya,
                "пришло",
                istochnik_poluchatelya,
                summa,
            ))

    rezultat = pd.DataFrame(obychnye_stroki, columns=KOLONKI)
    if not rezultat.empty:
        rezultat = rezultat.sort_values(
            ["Дата", "Фирма", "Сумма"], kind="stable", na_position="last"
        ).reset_index(drop=True)
    return rezultat, propusheno


def sprosit_tekst(podskazka, obyazatelno=True):
    while True:
        znachenie = input(podskazka).strip()
        if znachenie == "-1":
            return None
        if znachenie or not obyazatelno:
            return znachenie
        print("Поле не должно быть пустым")


def sprosit_ruchnuyu_stroku():
    while True:
        tekst_daty = sprosit_tekst("Введите дату в виде 01.08.2026 (-1 — назад): ")
        if tekst_daty is None:
            return None
        data = preobrazovat_datu(tekst_daty)
        if not pd.isna(data):
            break
        print("Не удалось прочитать дату. Введите её ещё раз.")

    firma = sprosit_tekst("Введите фирму (-1 — назад): ")
    if firma is None:
        return None
    kontragent = sprosit_tekst("Введите контрагента (-1 — назад): ")
    if kontragent is None:
        return None

    print("1 - пришло")
    print("2 - ушло")
    napravlenie = input("Введите номер (-1 — назад): ").strip()
    while napravlenie not in {"1", "2", "-1"}:
        napravlenie = input("Введите 1 или 2: ").strip()
    if napravlenie == "-1":
        return None
    napravlenie = "пришло" if napravlenie == "1" else "ушло"

    tip = sprosit_tekst("Введите тип операции (-1 — назад): ")
    if tip is None:
        return None
    naimenovanie = sprosit_tekst(
        "Введите наименование или нажмите Enter, если его нет: ",
        obyazatelno=False,
    )
    if naimenovanie is None:
        return None

    while True:
        tekst_summy = sprosit_tekst("Введите сумму (-1 — назад): ")
        if tekst_summy is None:
            return None
        summa = denezhnoe_chislo(tekst_summy)
        if summa is not None and summa > 0:
            break
        print("Введите положительное число, например 1250,50")

    kommentariy = sprosit_tekst(
        "Введите комментарий или нажмите Enter, если его нет: ",
        obyazatelno=False,
    )
    if kommentariy is None:
        return None
    return {
        "Дата": data,
        "Фирма": firma,
        "Контрагент": kontragent,
        "Пришло/ушло": napravlenie,
        "Тип операции": tip,
        "Наименование": naimenovanie,
        "Сумма": float(summa),
        "Комментарий": kommentariy,
    }


def nayti_tablicu_v_sessii(tablicy, imya_fayla):
    for tablica in tablicy:
        if os.path.abspath(tablica.imya) == os.path.abspath(imya_fayla):
            return tablica
    return None


def sohranit_rezultat(tablicy, imya_fayla, df):
    tablica = nayti_tablicu_v_sessii(tablicy, imya_fayla)
    novaya = tablica is None
    if tablica is None:
        tablica = Tablica(imya_fayla, df)
    else:
        tablica.df = df

    tablica.df.to_excel(tablica.imya, index=False)
    oformit_tablicu_vseh_operaciy(tablica.imya)
    print("Таблица всех операций сохранена:", tablica.imya)
    return [tablica] if novaya else []


def tablica_vseh_operaciy(tablicy, papka_mesaca):
    print("\nТаблица всех операций")
    print("1 - составить автоматически")
    print("2 - добавить одну запись вручную")
    print("-1 - вернуться в главное меню")

    vybor = input("Введите номер пункта: ").strip()
    while vybor not in {"1", "2", "-1"}:
        vybor = input("Нет такого номера. Введите ещё раз: ").strip()
    if vybor == "-1":
        return []

    imya_fayla = os.path.join(papka_mesaca, "таблица_всех_операций.xlsx")
    if vybor == "1":
        vybrannye = vybrat_istochniki(tablicy)
        if vybrannye is None:
            return []
        df, propusheno = sostavit_avtomaticheski(vybrannye)
        novye = sohranit_rezultat(tablicy, imya_fayla, df)
        print("Количество строк в таблице:", len(df))
        if propusheno:
            print(
                "Внимание: пропущено строк с неправильной суммой:",
                propusheno,
            )
        return novye

    ruchnaya_stroka = sprosit_ruchnuyu_stroku()
    if ruchnaya_stroka is None:
        return []

    tablica = nayti_tablicu_v_sessii(tablicy, imya_fayla)
    if tablica is not None:
        df = tablica.df.copy()
    elif os.path.isfile(imya_fayla):
        try:
            df = pd.read_excel(imya_fayla)
        except Exception as oshibka:
            print("Ошибка чтения таблицы всех операций:", oshibka)
            return []
    else:
        df = pd.DataFrame(columns=KOLONKI)

    df = df.reindex(columns=KOLONKI)
    df = pd.concat([df, pd.DataFrame([ruchnaya_stroka])], ignore_index=True)
    return sohranit_rezultat(tablicy, imya_fayla, df)
